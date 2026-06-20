from langchain_groq import ChatGroq
from langchain_core.prompts import PromptTemplate
from dotenv import load_dotenv
import json
import io
from jsonschema import ValidationError, validate
from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from tenacity import retry, stop_after_attempt, wait_exponential
from fastapi.middleware.cors import CORSMiddleware
from pypdf import PdfReader
from docx import Document
from striprtf.striprtf import rtf_to_text


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".rtf"}
SUPPORTED_DISPLAY = ", ".join(sorted(e.lstrip(".").upper() for e in SUPPORTED_EXTENSIONS))


origins = ["http://localhost:3000"]

app = FastAPI(title="AI Resume Analyzer",
            version="1.0.0")


app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

load_dotenv()

# Model
model = ChatGroq(
    model='llama-3.1-8b-instant',
    timeout=45,
    temperature=0.1
)

#  Schema
with open("json_schema.json", "r", encoding="utf-8") as f:
    schema = json.load(f)

# Prompt template 
template = PromptTemplate(
    template="""
You are an advanced,strict and accurate, enterprise-grade Applicant Tracking System (ATS) algorithm and an elite Technical Recruiter. Your task is to perform a deep-dive, ruthless, and highly actionable critique of the provided resume.

If a target Job Description (JD) is provided, evaluate the resume primarily against that specific role. If no JD is provided, evaluate it against general modern industry standards for the candidate's apparent domain.

### INPUT DATA:
- **Target Job Description:** 
{job_description}

- **Candidate Resume Raw Text:** 
{resume_text}

---

### CRITICAL OUTPUT CONSTRAINT:
- Return **ONLY** a raw, valid JSON object. 
- Do **NOT** wrap the response in markdown blocks (e.g., do not use ```json ... ```).
- Do **NOT** include any introductory text, pleasantries, or concluding notes.
- The output **MUST** strictly validate against this JSON Schema:
{schema}

---

### DETAILED DATA POPULATION INSTRUCTIONS:

1. **`ats_score` Evaluation:**
   - `overall`: A weighted average (0-100) based on the sub-scores below. Be realistic—do not inflate scores. An average resume should score around 60-70.
   - `grade`: Map the overall score to a letter grade: 90+ = A+, 85-89 = A, 80-84 = B+, 70-79 = B, 60-69 = C+, 50-59 = C, 40-49 = D, <40 = F,
      ONLY GIVE THE GRADES AS MENTIONED dont give like A- OR B- give as it is mentioned
   - `passing`: Boolean. Set to `true` if `overall` is 70 or above; otherwise `false`.
   - `breakdown`: 
     - `keyword_match`: Score (0-100) tracking the volume and importance of missing vs. matched skills.
     - `formatting`: Score (0-100). Deduct points for sloppy structure, missing essential sections, or clear parsing anomalies.
     - `section_completeness`: Score (0-100) verifying if the resume contains Contact, Summary, Experience, Education, and Skills.
     - `quantification`: Score (0-100). High scores require metrics, data, and outcomes (e.g., "$10k saved", "25% speed increase"). Deduct heavily if bullets are purely task-oriented.
     - `readability`: Score (0-100) judging wordiness, passive voice, or dense walls of text.

2. **`keyword_analysis` Evaluation:**
   - `matched`: Array of critical technical skills, methodologies, and tools present in both the resume and JD.
   - `missing`: Array of vital keywords/skills present in the JD that the candidate completely omitted. 
   - `overused`: Generic buzzwords or clichés used in the resume that add no value (e.g., "Synergy", "Hard worker", "Detail-oriented", "Go-getter").
   - `match_percentage`: Math calculation: `(matched / (matched + missing)) * 100`.

3. **`strengths` & `weaknesses` Arrays:**
   - Identify up to 3-5 structural or content-based elements.
   - For `strengths`, assign an `impact` rating ("high", "medium", or "low") based on how much it helps them pass an initial recruiter screen.
   - For `weaknesses`, assign a `severity` ("high", "medium", or "low") and provide a clear, constructive, actionable `fix`.
   - *Example sections:* "Work Experience", "Skills", "Header", "Summary", "Education".

4. **`section_feedback` Array:**
   - You must evaluate individual core sections: "Header/Contact Info", "Professional Summary", "Work Experience", "Education", and "Skills".
   - State if it is `present` (true/false), award a `score` (0-100) for that section's quality, and provide a concrete `suggestion` for enhancement.

5. **`bullet_rewrites` Array:**
   - Identify up to 3 weak, vague, or purely responsibility-based bullet points from the resume.
   - Provide the exact `original` text.
   - Provide a highly optimized `rewritten` version using the **Google X-Y-Z formula**: *Accomplished [X] as measured by [Y], by doing [Z]*. Ensure you inject metrics and stronger action verbs.
   - Give a compelling `reason` detailing why the rewrite is superior.

6. **`formatting_issues` Array:**
   - Scan for layout issues (e.g., inconsistent date formatting, lack of bullet points, mixing chronological order, over-crowded text, or missing critical contact information). Provide the specific `issue`, its `severity`, and an actionable `fix`.

7. **`meta` Information:**
   - Accurate counts for `word_count` and `page_count`.
   - Leave `file_name` and `analyzed_at` fields as placeholders or logical values; they will be handled programmatically.
   - `job_description_provided`: Set to `true` if a non-empty Job Description was provided; otherwise `false`.

Begin Analysis. Output pure JSON matching the schema definitions exactly:
""",
    input_variables=['resume_text', 'job_description'],
    partial_variables={"schema": json.dumps(schema)}
)



@retry(stop=stop_after_attempt(3), wait=wait_exponential(multiplier=1, min=2, max=10))
def call_llm(prompt):
    return model.invoke(prompt)



def get_extension(filename: str) -> str:
    """Returns lowercase extension including dot, e.g. '.pdf'. Empty string if no dot."""
    return ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""


def clean_llm_json(content: str) -> str:
    """Strips accidental markdown fences the model may emit."""
    content = content.strip()
    if content.startswith("```json"):
        content = content[7:]
    elif content.startswith("```"):
        content = content[3:]
    if content.endswith("```"):
        content = content[:-3]
    return content.strip()


# ── Text extractors ───────────────────────────────────────────────────────────
def _extract_pdf(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = [page.extract_text() for page in reader.pages if page.extract_text()]
    return "\n".join(pages).strip()


def _extract_docx(file_bytes: bytes) -> str:
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n".join(paragraphs).strip()


def _extract_plain(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "latin-1", "cp1252"):
        try:
            return file_bytes.decode(encoding).strip()
        except UnicodeDecodeError:
            continue
    raise ValueError("Unable to decode file as plain text.")


def _extract_rtf(file_bytes: bytes) -> str:
    raw = file_bytes.decode("latin-1", errors="ignore")
    return rtf_to_text(raw).strip()


EXTRACTORS = {
    ".pdf":  _extract_pdf,
    ".docx": _extract_docx,
    ".txt":  _extract_plain,
    ".md":   _extract_plain,
    ".rtf":  _extract_rtf,
}


def extract_text(file_bytes: bytes, filename: str) -> str:
    ext = get_extension(filename)
    
    extractor = EXTRACTORS.get(ext)
    
    if not extractor:
        
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported format '{ext or 'unknown'}'. Supported: {SUPPORTED_DISPLAY}"
        )
        
    try:
        return extractor(file_bytes)
    
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Failed to parse file: {str(e)}")



@app.post("/api/v1/analyzeResume")
async def analyze_resume(
    file: UploadFile = File(...),
    job_description: str = Form(default="")
):
    # Validate extension
    ext = get_extension(file.filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported format '{ext or 'unknown'}'. Supported formats: {SUPPORTED_DISPLAY}"
        )

    # Read and extract text
    file_bytes = await file.read()
    resume_text = extract_text(file_bytes, file.filename)

    if not resume_text:
        raise HTTPException(
            status_code=400,
            detail="File appears to be empty or unreadable."
        )

    # Build prompt and call LLM
    prompt = template.invoke({
        "resume_text": resume_text,
        "job_description": job_description if job_description else "None provided"
    })

    try:
        response = call_llm(prompt)
        content = clean_llm_json(response.content)
        data = json.loads(content)

        # Inject file-level metadata
        if "meta" in data:
            data["meta"]["file_name"] = file.filename
            data["meta"]["job_description_provided"] = bool(job_description)

        validate(instance=data,
                 schema=schema)
        
        return data

    except json.JSONDecodeError:
        
        raise HTTPException(status_code=500, detail="LLM returned invalid JSON syntax.")
    
    except ValidationError as e:
        
        raise HTTPException(status_code=500, detail=f"Schema validation failed: {e.message}")
    
    except Exception as e:
        
        raise HTTPException(status_code=500, detail=str(e))